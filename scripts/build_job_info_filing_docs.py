from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "legal_submission" / "filing"
SCREENS = ROOT / "docs" / "legal_submission" / "screens"
DATE = "2026년 8월 4일"
TEAL = "0F766E"
INK = "1F2937"
SUB = "64748B"
LINE = "CBD5E1"


def font(run, size=12, bold=False, color=INK):
    run.font.name = "Malgun Gothic"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def paragraph(doc, text="", size=12, bold=False, color=INK, after=5, align=WD_ALIGN_PARAGRAPH.LEFT):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.18
    font(p.add_run(text), size, bold, color)
    return p


def heading(doc, text, main=False):
    p = paragraph(doc, text, 16 if main else 14, True, "0F4C5C", 7)
    if not main:
        p.paragraph_format.space_before = Pt(10)
        p_pr = p._p.get_or_add_pPr()
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "5")
        bottom.set(qn("w:color"), TEAL)
        border.append(bottom)
        p_pr.append(border)
    return p


def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(cell):
    tc_pr = cell._tc.get_or_add_tcPr()
    box = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), LINE)
        box.append(el)
    tc_pr.append(box)


def table(doc, headers, rows, widths=None, font_size=10.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, text in enumerate(headers):
        cell = t.rows[0].cells[i]
        borders(cell); shade(cell, "E6FFFB")
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        if widths: cell.width = Cm(widths[i])
        font(cell.paragraphs[0].add_run(text), font_size, True, "134E4A")
    for row in rows:
        cells = t.add_row().cells
        for i, text in enumerate(row):
            borders(cells[i])
            if widths: cells[i].width = Cm(widths[i])
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.08
            font(p.add_run(str(text)), font_size, i == 0)
    paragraph(doc, "", after=1)
    return t


def bullets(doc, items):
    for text in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.line_spacing = 1.12
        font(p.add_run(text), 12)


def callout(doc, title, body, fill="F0FDFA"):
    t = doc.add_table(rows=1, cols=1)
    cell = t.cell(0, 0); borders(cell); shade(cell, fill)
    p = cell.paragraphs[0]; p.paragraph_format.space_after = Pt(3)
    font(p.add_run(title), 12, True, "0F4C5C")
    p2 = cell.add_paragraph(); p2.paragraph_format.space_after = Pt(0); p2.paragraph_format.line_spacing = 1.15
    font(p2.add_run(body), 12)
    paragraph(doc, "", after=1)


def setup(title):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21.0); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(1.55); sec.bottom_margin = Cm(1.55)
    sec.left_margin = Cm(1.75); sec.right_margin = Cm(1.75)
    normal = doc.styles["Normal"]
    normal.font.name = "Malgun Gothic"
    normal._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "맑은 고딕")
    normal.font.size = Pt(12)
    doc.core_properties.title = title
    doc.core_properties.author = "잇닿"
    return doc


def add_footer(doc):
    for sec in doc.sections:
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        font(p.add_run("잇닿 직업정보제공사업 신고 첨부자료"), 9, False, SUB)


def build_plan():
    doc = setup("잇닿 직업정보제공사업 사업계획서")
    heading(doc, "직업정보제공사업 사업계획서", True)
    paragraph(doc, "서비스명: 잇닿(Itdat) · 사업소 소재지: 경기도 수원시 권선구", 12, True, "334155")
    paragraph(doc, f"작성일: {DATE} · 신고인: ____________________", 10, False, SUB, 7)
    callout(doc, "사업 개요", "병원·의원·약국이 구인공고를 직접 게시하고, 간호사·간호조무사·약사·약국 전산/사무직 워커가 공고를 비교해 직접 지원하는 온라인 직업정보 제공 서비스입니다. 사업장은 지원자를 직접 검토하여 채용 여부를 결정합니다.")

    heading(doc, "1. 사업 목적과 제공 내용")
    table(doc, ["구분", "제공 내용"], [
        ["구인 사업장", "근무일·시간·장소·직무·자격요건·임금 조건을 직접 입력하고 지원자를 직접 검토"],
        ["구직 워커", "직군·지역에 맞는 공고를 검색·비교하고 원하는 공고에 직접 지원"],
        ["잇닿", "공고 게시, 지원상태, 자격서류 확인, 채팅, GPS·동적 QR 근태, 휴가·급여 검토용 SaaS 제공"],
    ], [3.0, 13.2])
    bullets(doc, [
        "잇닿은 특정 구인자와 구직자를 선정·추천·배정하지 않습니다.",
        "근로조건 협상과 근로계약은 병원·약국과 워커가 직접 진행합니다.",
        "임금은 병원·약국이 워커 계좌로 직접 지급하며 잇닿은 수취·보관·재지급하지 않습니다.",
    ])

    heading(doc, "2. 운영 및 관리 계획")
    table(doc, ["영역", "운영 방법"], [
        ["구인정보 관리", "사업장 정보와 공고 필수항목을 확인하고 허위·위법·차별 공고 신고 및 삭제 절차 운영"],
        ["구직자 보호", "최소정보만 수집하고 면허·이력서 등 민감 서류는 비공개 저장 및 권한 있는 관리자만 열람"],
        ["직군 검증", "간호사·간호조무사·약사 자격서류를 확인하며 약국 사무직 공고에는 면허업무를 포함하지 않도록 제한"],
        ["거래 구조", "사업장 직접 채용·직접 계약·직접 임금 지급 원칙을 화면과 약관에 고지"],
        ["사후 관리", "문의·신고 접수, 접근기록·근태 인증로그 보존, 변경사항 발생 시 변경신고 검토"],
    ], [3.0, 13.2])

    doc.add_page_break()
    heading(doc, "3. 수익모델")
    paragraph(doc, "수익은 병원·약국이 사용하는 업무용 소프트웨어의 월 정액 이용료와 관리자 좌석 등 부가 기능 이용료로 구성합니다. 워커 가입비, 소개비, 채용 성공수수료 및 임금 연동 수수료는 받지 않습니다.")
    table(doc, ["상품 예시", "월 이용료", "비고"], [
        ["Clinic", "59,000원", "소형 병·의원 채용·근태·휴가·급여 검토"],
        ["Basic / Pro", "79,000원 / 149,000원", "직원·공고·관리자·운영 자동화 범위에 따라 구분"],
        ["Pharmacy", "59,000원", "약사·전산직 운영, 관리자 좌석 추가 월 20,000원"],
        ["Pharmacy Plus", "99,000원", "관리자 3명, 대체약사 풀·반복 일정 운영"],
    ], [4.0, 3.2, 9.0])

    heading(doc, "4. 1차 연도 수입·지출 및 수지예산")
    paragraph(doc, "파일럿부터 유료 전환까지의 보수적 내부 예산이며 실제 가입 사업장 수와 사용량에 따라 달라질 수 있습니다.", 10, False, SUB)
    table(doc, ["구분", "산출 근거", "연간 금액"], [
        ["수입: SaaS 구독", "유료 사업장 평균 10곳 × 평균 79,000원 × 9개월", "7,110,000원"],
        ["수입: 부가 기능", "관리자 좌석·알림 등 선택 기능", "900,000원"],
        ["수입 합계", "", "8,010,000원"],
        ["지출: 서버·DB·도메인", "클라우드, 저장소, 모니터링", "2,400,000원"],
        ["지출: 알림·인증·API", "문자·푸시·지도·본인확인 준비", "900,000원"],
        ["지출: 마케팅·영업", "병원·약국 파일럿 자료와 방문", "1,600,000원"],
        ["지출: 워커 캠페인", "예산 한도형 프로필·첫 근무 리워드", "1,000,000원"],
        ["지출: 행정·법무·회계", "신고, 약관, 노무·세무 자문", "1,000,000원"],
        ["지출: 기타 운영", "고객지원·예비비", "600,000원"],
        ["지출 합계", "", "7,500,000원"],
        ["예상 수지", "수입 합계 − 지출 합계", "+510,000원"],
    ], [3.4, 8.5, 4.3], 9.5)
    callout(doc, "자금 및 확장 원칙", "초기 개발·운영비는 자기자금으로 충당하고, 파일럿 결과에 따라 지출합니다. 사업 범위가 적극적인 인재 추천·알선으로 변경되는 경우에는 운영 전에 유료직업소개사업 등록 필요성을 별도로 검토합니다.", "FFF7ED")
    add_footer(doc)
    path = OUT / "01_잇닿_직업정보제공사업_사업계획서.docx"
    doc.save(path)
    return path


def build_service():
    doc = setup("잇닿 직업정보제공사업 서비스 설명자료")
    heading(doc, "잇닿 서비스 설명자료", True)
    paragraph(doc, "직업정보제공사업 신규 신고 첨부자료", 14, True, "334155")
    paragraph(doc, f"작성일: {DATE} · 사업장 소재지: 경기도 수원시 권선구", 10, False, SUB)
    callout(doc, "한 문장 설명", "병원·의원·약국이 공고를 직접 게시하고 의료 워커가 직접 찾아 지원하며, 채용 이후에는 기존 직원과 단기근로자의 근태·휴가·급여 검토를 함께 관리하는 SaaS입니다.")

    heading(doc, "1. 이용자와 역할")
    table(doc, ["이용자", "직접 수행하는 일"], [
        ["병원·의원·약국", "공고 작성, 근무조건·임금 제시, 지원자 검토, 채용 결정, 근로계약, 임금 직접 지급"],
        ["워커", "공고 검색·비교, 지원 여부 결정, 직접 지원, 근로조건 확인, 출퇴근·휴가 신청"],
        ["잇닿", "직업정보 게시 공간과 지원상태·자격서류·채팅·근태·휴가·급여 검토 소프트웨어 제공"],
    ], [3.2, 13.0])

    heading(doc, "2. 실제 이용 흐름")
    table(doc, ["단계", "내용", "결정 주체"], [
        ["1", "사업장이 날짜·시간·업무·자격·임금을 입력해 공고 게시", "사업장"],
        ["2", "워커가 지역·직군에 맞는 여러 공고를 비교", "워커"],
        ["3", "워커가 원하는 공고에 직접 지원", "워커"],
        ["4", "사업장이 지원자 서류를 확인하고 수락 또는 거절", "사업장"],
        ["5", "근로조건 확인과 근로계약 체결", "사업장·워커"],
        ["6", "GPS 우선·동적 QR 보완으로 출퇴근 기록", "사업장 정책"],
        ["7", "근태와 예상금액 확인 후 워커 계좌로 직접 지급", "사업장"],
    ], [1.5, 10.8, 3.9], 10)
    callout(doc, "직업정보제공 원칙", "조건 기반 검색·거리순 표시·직군 알림은 정보 탐색을 돕는 기능입니다. 잇닿이 지원자를 평가해 순위를 정하거나 특정 사업장·워커를 채용 대상으로 추천하지 않습니다.", "EFF6FF")

    heading(doc, "3. 지원 직군과 사업장")
    table(doc, ["사업장", "직군", "주요 정보"], [
        ["병원·의원", "간호사, 간호조무사", "면허·자격, 부서, 근무일시, 시급, 담당업무"],
        ["약국", "약사, 약국 전산·사무직", "약사면허 또는 이력서, 전산 프로그램, 처방량, 인수인계, 면허업무 구분"],
    ], [3.0, 4.5, 8.7])

    doc.add_page_break()
    heading(doc, "4. 화면으로 보는 운영 구조")
    table(doc, ["관리자 공고 등록", "병원이 근무조건을 직접 작성"], [["입력 항목", "근무일·시간·자격·부서·시급·업무·안내사항"], ["게시 결정", "병원·약국 관리자"], ["잇닿 역할", "입력 화면과 게시 기능 제공"]], [4.0, 12.2])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SCREENS / "admin_shift_new.png"), width=Cm(6.6))
    paragraph(doc, "그림 1. 사업장이 필요한 직군·일정·임금을 직접 입력하는 공고 등록 화면", 10, False, SUB, 6, WD_ALIGN_PARAGRAPH.CENTER)

    heading(doc, "5. 근태·휴가·급여 검토 SaaS")
    bullets(doc, [
        "신규 채용 인력과 기존 직원을 하나의 직원관리·월 근태 화면에서 관리합니다.",
        "출퇴근은 GPS 원터치를 우선 사용하고 실내 위치 오류 시 60초 동적 QR로 보완합니다.",
        "근태기록은 지각·조퇴·결근·휴가·총 근무시간을 집계해 사업장이 확인합니다.",
        "급여 화면의 금액은 검토용 예상금액이며 최종 계산·세금·보험·지급 책임은 사업장에 있습니다.",
    ])
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(SCREENS / "admin_payroll.png"), width=Cm(6.6))
    paragraph(doc, "그림 2. 근태와 예상금액을 확인하고 사업장이 직접 지급 결과를 기록하는 화면", 10, False, SUB, 6, WD_ALIGN_PARAGRAPH.CENTER)

    heading(doc, "6. 수익 구조")
    table(doc, ["항목", "운영 여부", "설명"], [
        ["사업장 월 정액 SaaS", "운영", "공고·직원·근태·휴가·급여 검토·관리자 기능 이용료"],
        ["관리자 좌석 등 부가기능", "운영", "소프트웨어 사용량에 따른 선택 이용료"],
        ["워커 가입비·소개비", "미운영", "워커에게 청구하지 않음"],
        ["채용 성공수수료", "미운영", "지원·수락·채용 여부와 이용료를 연동하지 않음"],
        ["임금 보관·지급대행", "미운영", "사업장이 워커에게 직접 지급"],
    ], [4.0, 2.5, 9.7])

    heading(doc, "7. 개인정보·허위정보 보호")
    bullets(doc, [
        "구인 사업장과 사용자 계정을 확인하고 허위공고 신고·차단·삭제 절차를 운영합니다.",
        "면허·이력서·계좌정보는 공개 공고에 표시하지 않고 권한이 있는 사용자만 접근합니다.",
        "QR 토큰은 짧은 시간만 유효하며 원문 대신 해시와 인증 시도 로그를 서버에 저장합니다.",
        "개인정보 처리방침, 이용약관, 위치정보 동의와 마케팅 선택동의를 구분합니다.",
    ])

    heading(doc, "8. 잇닿이 하지 않는 일")
    bullets(doc, [
        "특정 구인자와 구직자를 선별해 개별적으로 추천·배정하거나 채용을 권유하는 행위",
        "근로조건과 임금을 대신 협상하거나 근로계약을 대리·체결하는 행위",
        "워커 임금을 수취·보관·재지급하거나 임금에서 수수료를 공제하는 행위",
        "약국 전산·사무직에게 조제·복약지도 등 약사 면허업무를 수행하게 하는 공고 제공",
    ])
    callout(doc, "향후 변경 시", "적극적인 인재 추천·알선, 성공보수 또는 지급대행 등으로 사업모델을 변경하기 전에는 관할기관에 변경신고 및 유료직업소개사업 등록 필요성을 확인하겠습니다.", "FFF7ED")
    add_footer(doc)
    path = OUT / "02_잇닿_직업정보제공사업_서비스설명자료.docx"
    doc.save(path)
    return path


if __name__ == "__main__":
    OUT.mkdir(parents=True, exist_ok=True)
    print(build_plan())
    print(build_service())
